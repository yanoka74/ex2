import pytest
import openpyxl
from pathlib import Path


COLUMNS = [
    "고객 주문 마켓별칭",
    "마켓명",
    "주문일",
    "주문번호",
    "판매자 고유코드",
    "상품코드",
    "상품명",
    "옵션명",
    "주문수량",
    "주문가격",
    "수령자",
    "전화번호",
    "주소",
]


def create_test_xlsx(rows: list[dict], path: Path) -> Path:
    """리스트 of dict를 받아 .xlsx 파일을 생성한다.

    각 dict의 키는 COLUMNS 중 일부 또는 전부.
    누락된 컬럼은 빈 문자열로 채운다.
    """
    wb = openpyxl.Workbook()
    ws = wb.active

    # 헤더 행 작성
    for col_idx, col_name in enumerate(COLUMNS, start=1):
        ws.cell(row=1, column=col_idx, value=col_name)

    # 데이터 행 작성
    for row_idx, row_data in enumerate(rows, start=2):
        for col_idx, col_name in enumerate(COLUMNS, start=1):
            value = row_data.get(col_name, "")
            ws.cell(row=row_idx, column=col_idx, value=value)

    wb.save(path)
    return path


def create_test_xlsx_with_columns(columns: list[str], rows: list[list], path: Path) -> Path:
    """커스텀 컬럼 헤더로 .xlsx 파일을 생성한다. 필수 컬럼 누락 테스트용."""
    wb = openpyxl.Workbook()
    ws = wb.active

    for col_idx, col_name in enumerate(columns, start=1):
        ws.cell(row=1, column=col_idx, value=col_name)

    for row_idx, row_data in enumerate(rows, start=2):
        for col_idx, value in enumerate(row_data, start=1):
            ws.cell(row=row_idx, column=col_idx, value=value)

    wb.save(path)
    return path


@pytest.fixture
def xlsx_factory(tmp_path):
    """테스트 엑셀 파일을 동적으로 생성하는 팩토리 픽스처."""
    def _create(rows: list[dict], filename: str = "test.xlsx") -> Path:
        return create_test_xlsx(rows, tmp_path / filename)
    return _create


@pytest.fixture
def xlsx_custom_columns_factory(tmp_path):
    """커스텀 컬럼으로 테스트 엑셀 파일을 생성하는 팩토리 픽스처."""
    def _create(columns: list[str], rows: list[list], filename: str = "test.xlsx") -> Path:
        return create_test_xlsx_with_columns(columns, rows, tmp_path / filename)
    return _create
